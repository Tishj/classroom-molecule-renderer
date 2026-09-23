"""Notebook-friendly Word export. python-docx is an optional dependency."""

from dataclasses import dataclass, replace
from io import BytesIO
from math import isfinite

from .renderer import MoleculeError, RenderOptions, _figure


@dataclass(frozen=True)
class MoleculeEntry:
    name: str
    smiles: str
    stereochemistry: str | None = None


@dataclass(frozen=True)
class DocumentOptions:
    title: str = "Structuurformules"
    include_names: bool = True
    numbered: bool = False
    scale: float = 1.0

    def __post_init__(self):
        if not isfinite(self.scale) or not 0.1 <= self.scale <= 3:
            raise ValueError("Document scale must be between 0.1 and 3.")


def parse_molecule_lines(text):
    """Read one 'name | SMILES' or bare SMILES per line; ignore empty lines."""
    entries = []
    for line_number, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        fields = line.split("|")
        if len(fields) == 1:
            name, smiles = "", fields[0].strip()
        elif len(fields) == 2:
            name, smiles = (field.strip() for field in fields)
        else:
            raise ValueError(f"Line {line_number}: use name | SMILES (one separator).")
        if not smiles:
            raise ValueError(f"Line {line_number}: SMILES is missing.")
        entries.append(MoleculeEntry(name, smiles))
    if not entries:
        raise ValueError("Add at least one molecule.")
    if len(entries) > 100:
        raise ValueError("Export at most 100 molecules at a time.")
    return entries


@dataclass(frozen=True)
class _MoleculeImage:
    entry: MoleculeEntry
    png: bytes
    width_pt: float
    height_pt: float


def _prepare_images(entries, options=None):
    entries = list(entries)
    if not entries or len(entries) > 100:
        raise ValueError("Provide between 1 and 100 molecules.")
    images = []
    for index, entry in enumerate(entries, 1):
        try:
            entry_options = options or RenderOptions()
            if entry.stereochemistry is not None:
                entry_options = replace(entry_options, stereochemistry=entry.stereochemistry)
            fig = _figure(entry.smiles, entry_options)
        except ValueError as error:
            raise MoleculeError(f"Molecule {index} ({entry.name or entry.smiles}): {error}") from error
        stream = BytesIO()
        fig.savefig(stream, format="png", dpi=300, facecolor="white")
        width, height = fig.get_size_inches() * 72
        images.append(_MoleculeImage(entry, stream.getvalue(), float(width), float(height)))
    return images


def _document_bytes(images, document_options=None):
    # Kept separate from rendering so document construction can reuse images.
    try:
        from docx import Document
        from docx.shared import Mm, Pt, RGBColor
    except ImportError as error:
        raise ImportError("Word export needs python-docx: pip install -e '.[docx]'") from error

    settings = document_options or DocumentOptions()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)
    for style_name in ("Normal", "Title"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Title"].font.size = Pt(22)
    # Some python-docx templates carry a decorative Title border.
    for style_name in ("Normal", "Title"):
        for border in doc.styles[style_name].element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.core_properties.title = settings.title
    if settings.title.strip():
        title = doc.add_paragraph(settings.title, style="Title")
        title.paragraph_format.space_after = Pt(16)
        title.paragraph_format.keep_with_next = True

    # One common scale preserves the same font size across every molecule.
    # Reserve room for a caption, even for tall RDKit coordinate layouts.
    max_width = (section.page_width - section.left_margin - section.right_margin) / 12700
    max_height = (section.page_height - section.top_margin - section.bottom_margin) / 12700 - 90
    scale = min(settings.scale, max_width / max(image.width_pt for image in images),
                max_height / max(image.height_pt for image in images))
    for index, image in enumerate(images, 1):
        caption = image.entry.name if settings.include_names else ""
        if settings.numbered:
            caption = f"{index}. {caption}".rstrip()
        if caption:
            paragraph = doc.add_paragraph()
            paragraph.add_run(caption).bold = True
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(2)
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_after = Pt(10)
        picture = paragraph.add_run().add_picture(BytesIO(image.png),
                         width=Pt(image.width_pt * scale), height=Pt(image.height_pt * scale))
        picture._inline.docPr.set("descr", f"Structural formula: {image.entry.smiles}")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def to_docx(entries, options=None, document_options=None):
    """Return DOCX bytes containing labeled molecule images at a common scale.

    Entries are MoleculeEntry objects, in document order. An empty title removes
    the heading. Molecules are high-resolution images, not editable Word shapes.
    All entries are validated before producing a document; none are skipped.
    """
    return _document_bytes(_prepare_images(entries, options), document_options)
