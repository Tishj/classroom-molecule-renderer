from pathlib import Path
import re

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "quiz_assets"
OUT = ROOT / "HAVO_scheikunde_namen_en_formules_syllabus_2027.docx"
ASSETS.mkdir(exist_ok=True)

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GREY = "F4F6F7"
MID_GREY = "D9D9D9"
TEXT = "1F1F1F"
ACCENT = "0E6B78"


def mol(labels, orders=None, branches=None, text=None):
    return {"labels": labels, "orders": orders or [1] * max(0, len(labels) - 1),
            "branches": branches or {}, "text": text}


def draw_molecule(ax, spec, fontsize=16):
    ax.set_axis_off()
    if spec.get("text"):
        ax.text(0.5, 0.5, spec["text"], ha="center", va="center",
                fontsize=fontsize, fontfamily="DejaVu Sans", color="#111111",
                transform=ax.transAxes)
        return
    labels = spec["labels"]
    n = len(labels)
    spacing = 1.55
    xs = [(i - (n - 1) / 2) * spacing for i in range(n)]
    ys = [0.0] * n
    for i, order in enumerate(spec["orders"]):
        x1, x2 = xs[i] + 0.55, xs[i + 1] - 0.55
        if order == 1:
            ax.plot([x1, x2], [0, 0], lw=1.9, color="#111111")
        else:
            ax.plot([x1, x2], [0.055, 0.055], lw=1.7, color="#111111")
            ax.plot([x1, x2], [-0.055, -0.055], lw=1.7, color="#111111")
    for i, label in enumerate(labels):
        if label == "OH":
            ax.text(xs[i], ys[i], "O", ha="center", va="center", fontsize=fontsize,
                    fontfamily="DejaVu Sans", color="#111111",
                    bbox=dict(facecolor="white", edgecolor="none", pad=0.4))
            ax.text(xs[i] + 0.22, ys[i], "H", ha="center", va="center", fontsize=fontsize,
                    fontfamily="DejaVu Sans", color="#111111")
        elif label == "HO":
            ax.text(xs[i], ys[i], "O", ha="center", va="center", fontsize=fontsize,
                    fontfamily="DejaVu Sans", color="#111111",
                    bbox=dict(facecolor="white", edgecolor="none", pad=0.4))
            ax.text(xs[i] - 0.22, ys[i], "H", ha="center", va="center", fontsize=fontsize,
                    fontfamily="DejaVu Sans", color="#111111")
        else:
            ax.text(xs[i], ys[i], label, ha="center", va="center", fontsize=fontsize,
                    fontfamily="DejaVu Sans", color="#111111",
                    bbox=dict(facecolor="white", edgecolor="none", pad=0.6))
    for idx, items in spec.get("branches", {}).items():
        for branch in items:
            label, direction, order = branch
            dx, dy = (0, 0.90) if direction == "up" else (0, -0.90)
            bx, by = xs[idx] + dx, dy
            if order == 1:
                ax.plot([xs[idx], bx], [0.27 * (1 if dy > 0 else -1), by - 0.27 * (1 if dy > 0 else -1)],
                        lw=1.9, color="#111111")
            else:
                off = 0.05
                ax.plot([xs[idx] - off, bx - off], [0.27 * (1 if dy > 0 else -1), by - 0.27 * (1 if dy > 0 else -1)],
                        lw=1.7, color="#111111")
                ax.plot([xs[idx] + off, bx + off], [0.27 * (1 if dy > 0 else -1), by - 0.27 * (1 if dy > 0 else -1)],
                        lw=1.7, color="#111111")
            if label == "OH":
                ax.text(bx, by, "O", ha="center", va="center", fontsize=fontsize,
                        fontfamily="DejaVu Sans", color="#111111",
                        bbox=dict(facecolor="white", edgecolor="none", pad=0.4))
                ax.text(bx + 0.22, by, "H", ha="center", va="center", fontsize=fontsize,
                        fontfamily="DejaVu Sans", color="#111111")
            else:
                ax.text(bx, by, label, ha="center", va="center", fontsize=fontsize,
                        fontfamily="DejaVu Sans", color="#111111",
                        bbox=dict(facecolor="white", edgecolor="none", pad=0.6))
    ymax = 1.45 if spec.get("branches") else 0.8
    ax.set_xlim(min(xs) - 1.05, max(xs) + 1.05)
    ax.set_ylim(-ymax, ymax)


def structure_choices(filename, choices):
    path = ASSETS / filename
    fig, axes = plt.subplots(2, 2, figsize=(10, 4.2), dpi=220)
    fig.patch.set_facecolor("white")
    for ax, label, spec in zip(axes.flat, "ABCD", choices):
        ax.add_patch(FancyBboxPatch((0.01, 0.05), 0.98, 0.90,
                                   boxstyle="round,pad=0.012,rounding_size=0.025",
                                   transform=ax.transAxes, facecolor="white",
                                   edgecolor="#CBD5E1", linewidth=1.1))
        ax.text(0.055, 0.82, label, transform=ax.transAxes, fontsize=13,
                fontweight="bold", color="#17365D")
        draw_molecule(ax, spec, fontsize=15)
    plt.subplots_adjust(left=0.02, right=0.98, top=0.98, bottom=0.02, wspace=0.08, hspace=0.10)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def structure_prompt(filename, spec):
    path = ASSETS / filename
    fig, ax = plt.subplots(figsize=(9, 1.45), dpi=220)
    fig.patch.set_facecolor("white")
    ax.add_patch(FancyBboxPatch((0.01, 0.08), 0.98, 0.84,
                               boxstyle="round,pad=0.012,rounding_size=0.025",
                               transform=ax.transAxes, facecolor="white",
                               edgecolor="#CBD5E1", linewidth=1.1))
    draw_molecule(ax, spec, fontsize=17)
    plt.subplots_adjust(left=0.02, right=0.98, top=0.96, bottom=0.04)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


# Syllabus-conforme structuurformules. Carbonylgroepen zijn steeds expliciet als C=O getekend.
qimg = {}
qimg[6] = structure_choices("q06.png", [
    mol(["CH₃", "CH₂", "CH₂", "OH"]),
    mol(["CH₃", "CH", "CH₃"], branches={1: [("OH", "up", 1)]}),
    mol(["CH₂", "CH", "CH₃"], orders=[2, 1]),
    mol(["CH₃", "CH₂", "C", "OH"], branches={2: [("O", "up", 2)]}),
])
qimg[7] = structure_prompt("q07.png", mol(["CH₃", "CH₂", "CH₂", "OH"]))
qimg[9] = structure_choices("q09.png", [
    mol(["CH₃", "C", "OH"], branches={1: [("O", "up", 2)]}),
    mol(["HO", "CH₂", "CH₂", "OH"]),
    mol(["CH₃", "CH₂", "OH"]),
    mol(["H", "C", "OH"], branches={1: [("O", "up", 2)]}),
])
qimg[15] = structure_choices("q15.png", [
    mol(["CH₃", "C", "O", "CH₂", "CH₃"], branches={1: [("O", "up", 2)]}),
    mol(["CH₃", "CH₂", "C", "O", "CH₃"], branches={2: [("O", "up", 2)]}),
    mol(["CH₃", "CH₂", "O", "CH₂", "CH₃"]),
    mol(["CH₃", "CH₂", "CH₂", "C", "OH"], branches={3: [("O", "up", 2)]}),
])
qimg[16] = structure_prompt("q16.png", mol(["CH₃", "CH", "CH₃"], branches={1: [("OH", "up", 1)]}))
qimg[17] = structure_choices("q17.png", [
    mol(["CH₂", "CH₂"], orders=[2]),
    mol(["CH₂", "CH", "CH₃"], orders=[2, 1]),
    mol(["CH₃", "CH₂", "CH₃"]),
    mol(["CH₂", "CH", "CH₂", "CH₃"], orders=[2, 1, 1]),
])
qimg[22] = structure_choices("q22.png", [
    mol(["CH₂", "CH", "CH₂", "CH₂", "CH₃"], orders=[2, 1, 1, 1]),
    mol(["CH₃", "CH", "CH", "CH₂", "CH₃"], orders=[1, 2, 1, 1]),
    mol(["CH₃", "CH₂", "CH₂", "CH₂", "CH₃"]),
    mol(["CH₂", "CH", "CH", "CH", "CH₃"], orders=[2, 1, 2, 1]),
])
qimg[26] = structure_choices("q26.png", [
    mol(["CH₃", "CH₂", "CH", "CH₂"], orders=[1, 1, 2]),
    mol(["CH₂", "CH", "CH", "CH₂"], orders=[2, 1, 2]),
    mol(["CH₂", "C", "CH", "CH₃"], orders=[2, 2, 1]),
    mol(["CH₃", "CH", "CH", "CH₃"], orders=[1, 2, 1]),
])
qimg[27] = structure_prompt("q27.png", mol(["CH₂", "CH", "CH₂", "CH₂", "CH₂", "CH₃"], orders=[2, 1, 1, 1, 1]))
qimg[35] = structure_choices("q35.png", [
    mol(["CH₃", "CH₂", "CH₂", "CH₂", "OH"]),
    mol(["CH₃", "CH", "CH₂", "CH₃"], branches={1: [("OH", "up", 1)]}),
    mol(["CH₃", "CH₂", "CH₂", "C", "OH"], branches={3: [("O", "up", 2)]}),
    mol(["CH₂", "CH", "CH₂", "CH₃"], orders=[2, 1, 1]),
])
qimg[36] = structure_prompt("q36.png", mol(["CH₃", "CH₂", "CH₂", "CH₂", "C", "OH"], branches={4: [("O", "up", 2)]}))
qimg[40] = structure_choices("q40.png", [
    mol(["HO", "CH₂", "CH₂", "OH"]),
    mol(["CH₃", "CH", "CH₂", "OH"], branches={1: [("OH", "up", 1)]}),
    mol(["HO", "CH₂", "CH", "CH₂", "OH"], branches={2: [("OH", "up", 1)]}),
    mol(["CH₃", "CH₂", "CH₂", "OH"]),
])
qimg[43] = structure_choices("q43.png", [
    mol(["CH₃", "CH₂", "OH"]),
    mol(["CH₃", "O", "CH₃"]),
    mol(["CH₃", "C", "OH"], branches={1: [("O", "up", 2)]}),
    mol(["CH₂", "CH", "OH"], orders=[2, 1]),
])
qimg[44] = structure_prompt("q44.png", mol(["CH₃", "C", "O", "CH₂", "CH₃"], branches={1: [("O", "up", 2)]}))
qimg[50] = structure_choices("q50.png", [
    mol([], text="—CH₂—CH(CH₃)—CH₂—CH(CH₃)—CH₂—CH(CH₃)—"),
    mol([], text="—CH₂—CH₂—CH₂—CH₂—CH₂—CH₂—"),
    mol([], text="—CH₂—CH=CH—CH₂—CH=CH—"),
    mol([], text="—CH₂—CH₂—O—CH₂—CH₂—O—"),
])


questions = [
 (1,"M","Wat is de molecuulformule van koolstofdioxide?",[("CO",1),("C2O",1),("CO2",1),("CO3",1)],"C"),
 (2,"M","Welke naam hoort bij de molecuulformule [[CO]]?",[("koolstofdioxide",0),("koolstofmono-oxide",0),("methaan",0),("zuurstof",0)],"B"),
 (3,"V","Welke naam hoort bij de verhoudingsformule [[Fe2O3]]?",[("ijzer(II)oxide",0),("ijzerchloride",0),("ijzeroxide-ion",0),("ijzer(III)oxide",0)],"D"),
 (4,"V","Welke verhoudingsformule hoort bij ijzer(III)oxide?",[("Fe2O3",1),("FeO",1),("Fe3O2",1),("FeO2",1)],"A"),
 (5,"M","Wat is de molecuulformule van methaan?",[("C2H6",1),("CH4",1),("CH3",1),("C2H4",1)],"B"),
 (6,"S","Welke structuurformule hoort bij propaan-2-ol?",None,"B"),
 (7,"S","Welke naam hoort bij deze structuurformule?",[("propaan-1-ol",0),("propaan-2-ol",0),("propaanzuur",0),("propeen",0)],"A"),
 (8,"M","Wat is de molecuulformule van water?",[("HO",1),("H2O2",1),("H2O",1),("OH",1)],"C"),
 (9,"S","Welke structuurformule hoort bij ethaanzuur?",None,"A"),
 (10,"V","Welke verhoudingsformule ontstaat uit [[Zn^2+]]- en [[Cl^-]]-ionen?",[("Zn2Cl",1),("ZnCl",1),("Zn2Cl3",1),("ZnCl2",1)],"D"),

 (11,"V","Welke naam hoort bij [[FeO]]?",[("ijzer(II)oxide",0),("ijzer(III)oxide",0),("ijzer(I)oxide",0),("ijzerhydroxide",0)],"A"),
 (12,"V","Wat is de verhoudingsformule van ijzer(II)chloride?",[("FeCl",1),("Fe2Cl",1),("FeCl3",1),("FeCl2",1)],"D"),
 (13,"M","Welke naam hoort bij [[HBr]]?",[("broom",0),("waterstofbromide",0),("bromide-ion",0),("waterstofperoxide",0)],"B"),
 (14,"M","Wat is de molecuulformule van methanol?",[("C2H6O",1),("CH2O",1),("CH4O",1),("C2H4O",1)],"C"),
 (15,"S","Welke structuurformule hoort bij ethylethanoaat?",None,"A"),
 (16,"S","Welke naam hoort bij deze structuurformule?",[("propaan-1-ol",0),("propeen",0),("propaanzuur",0),("propaan-2-ol",0)],"D"),
 (17,"S","Welke structuurformule heeft het monomeer waaruit polypropeen wordt gevormd?",None,"B"),
 (18,"M","Wat is de molecuulformule van propaan-2-ol?",[("C3H6O",1),("C3H8O",1),("C3H8O2",1),("C2H6O",1)],"B"),
 (19,"V","Welke verhoudingsformule ontstaat uit [[Mg^2+]]- en [[NO3^-]]-ionen?",[("MgNO3",1),("Mg2NO3",1),("Mg(NO3)2",1),("Mg2(NO3)3",1)],"C"),
 (20,"M","Welke naam hoort bij [[HCl]]?",[("waterstofchloride",0),("chloor",0),("chloride-ion",0),("waterstofhypochloriet",0)],"A"),

 (21,"M","Wat is de molecuulformule van stikstofdioxide?",[("NO",1),("N2O",1),("NO2",1),("N2O3",1)],"C"),
 (22,"S","Welke structuurformule hoort bij pent-2-een?",None,"B"),
 (23,"V","Welke naam hoort bij [[NaCl]]?",[("natriumchloride",0),("natriumchloraat",0),("natriumoxide",0),("natriumchloriet",0)],"A"),
 (24,"V","Welke verhoudingsformule ontstaat uit [[Al^3+]]- en [[O^2-]]-ionen?",[("AlO",1),("AlO2",1),("Al3O2",1),("Al2O3",1)],"D"),
 (25,"V","Wat is de verhoudingsformule van zinkchloride?",[("Zn2Cl",1),("ZnCl",1),("Zn2Cl3",1),("ZnCl2",1)],"D"),
 (26,"S","Welke structuurformule hoort bij buta-1,3-dieen?",None,"B"),
 (27,"S","Welke naam hoort bij deze structuurformule?",[("hex-1-een",0),("hexaan",0),("hex-2-een",0),("pent-1-een",0)],"A"),
 (28,"M","Wat is de molecuulformule van zwaveldioxide?",[("SO",1),("SO3",1),("SO2",1),("S2O",1)],"C"),
 (29,"M","Welke naam hoort bij [[H2O2]]?",[("water",0),("waterstofperoxide",0),("waterstofdioxide",0),("hydroxide",0)],"B"),
 (30,"V","Welke verhoudingsformule ontstaat uit [[Ni^2+]]- en [[Cl^-]]-ionen?",[("Ni2Cl",1),("NiCl",1),("Ni2Cl3",1),("NiCl2",1)],"D"),

 (31,"M","Wat is de molecuulformule van hexaan?",[("C6H14",1),("C6H12",1),("C5H12",1),("C6H6",1)],"A"),
 (32,"M","Welke naam hoort bij [[C5H12]]?",[("butaan",0),("penteen",0),("pentaan",0),("hexaan",0)],"C"),
 (33,"M","Welke naam hoort bij [[NH3]]?",[("ammonium",0),("ammoniak",0),("stikstofmono-oxide",0),("nitraat",0)],"B"),
 (34,"M","Wat is de molecuulformule van azijnzuur oftewel ethaanzuur?",[("CH2O2",1),("C2H6O",1),("C2H4O",1),("C2H4O2",1)],"D"),
 (35,"S","Welke structuurformule hoort bij butaan-1-ol?",None,"A"),
 (36,"S","Welke naam hoort bij deze structuurformule?",[("pentaan",0),("pentaan-1-ol",0),("pentaanzuur",0),("butaanzuur",0)],"C"),
 (37,"V","Welke naam hoort bij [[NaOH]]?",[("natriumoxide",0),("natriumwaterstof",0),("natriumperoxide",0),("natriumhydroxide",0)],"D"),
 (38,"V","Welke verhoudingsformule ontstaat uit [[Ca^2+]]- en [[CO3^2-]]-ionen?",[("Ca2CO3",1),("CaCO3",1),("Ca(CO3)2",1),("Ca3CO2",1)],"B"),
 (39,"M","Wat is de molecuulformule van zwaveltrioxide?",[("SO3",1),("SO2",1),("S3O",1),("S2O3",1)],"A"),
 (40,"S","Welke structuurformule hoort bij propaan-1,2,3-triol?",None,"C"),

 (41,"M","Wat is de molecuulformule van waterstofperoxide?",[("H2O",1),("HO2",1),("H3O",1),("H2O2",1)],"D"),
 (42,"M","Welke naam hoort bij [[C6H12O6]]?",[("sacharose",0),("glucose",0),("azijnzuur",0),("ammoniak",0)],"B"),
 (43,"S","Welke structuurformule hoort bij ethanol?",None,"A"),
 (44,"S","Welke naam hoort bij deze structuurformule?",[("methylethanoaat",0),("butaanzuur",0),("ethylethanoaat",0),("propan-1-ol",0)],"C"),
 (45,"V","Welke naam hoort bij [[NH4HCO3]]?",[("ammoniumcarbonaat",0),("ammoniumnitraat",0),("natriumwaterstofcarbonaat",0),("ammoniumwaterstofcarbonaat",0)],"D"),
 (46,"V","Welke verhoudingsformule ontstaat uit [[NH4^+]]- en [[SO4^2-]]-ionen?",[("NH4SO4",1),("(NH4)2SO4",1),("NH4(SO4)2",1),("(NH4)2SO3",1)],"B"),
 (47,"M","Wat is de molecuulformule van fosforzuur?",[("H3PO4",1),("H2PO4",1),("HPO3",1),("H3PO3",1)],"A"),
 (48,"V","Welke naam hoort bij [[CoO2]] als zuurstof aanwezig is als oxide-ion?",[("kobalt(I)oxide",0),("kobalt(II)oxide",0),("kobalt(IV)oxide",0),("kobalt(VI)oxide",0)],"C"),
 (49,"V","Welke verhoudingsformule ontstaat uit [[K^+]]- en [[PO4^3-]]-ionen?",[("KPO4",1),("K2PO4",1),("K2(PO4)3",1),("K3PO4",1)],"D"),
 (50,"S","Welke structuurformule stelt een fragment van polyetheen met drie monomeereenheden voor?",None,"B"),
]

years = {1:"2026", 11:"2025", 21:"2023", 31:"2022", 41:"2021"}
type_names = {"M":"Molecuulformule", "V":"Verhoudingsformule", "S":"Structuurformule"}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def add_formula(paragraph, formula, bold=False, size=None, color=TEXT):
    if "^" in formula:
        base, charge = formula.split("^", 1)
    else:
        base, charge = formula, ""
    for part in re.split(r"(\d+)", base):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.font.name = "Aptos"
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        if size: run.font.size = Pt(size)
        if part.isdigit(): run.font.subscript = True
    if charge:
        run = paragraph.add_run(charge)
        run.font.name = "Aptos"; run.font.superscript = True; run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        if size: run.font.size = Pt(size)


def add_marked(paragraph, text, bold=False, size=None):
    pos = 0
    for m in re.finditer(r"\[\[(.*?)\]\]", text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()]); r.bold = bold
            if size: r.font.size = Pt(size)
        add_formula(paragraph, m.group(1), bold=bold, size=size)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:]); r.bold = bold
        if size: r.font.size = Pt(size)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rPr.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); rPr.append(underline)
    new_run.append(rPr); t = OxmlElement("w:t"); t.text = text; new_run.append(t); hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_option(paragraph, letter, text, is_formula):
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.space_after = Pt(2.5)
    r = paragraph.add_run(f"{letter}.  "); r.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
    if is_formula:
        add_formula(paragraph, text, size=10.5)
    else:
        paragraph.add_run(text)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5); section.page_height = Inches(11)
section.top_margin = Inches(0.68); section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.75); section.right_margin = Inches(0.75)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(5)
for sty_name, size in (("Title",28),("Heading 1",19),("Heading 2",14)):
    sty = styles[sty_name]
    sty.font.name = "Aptos Display" if sty_name != "Heading 2" else "Aptos"
    sty.font.size = Pt(size); sty.font.color.rgb = RGBColor(0,0,0); sty.font.bold = True
    sty.paragraph_format.space_before = Pt(10); sty.paragraph_format.space_after = Pt(8)

# Neutraliseer de ingebouwde Word-rand van de Title-stijl.
title_ppr = styles["Title"]._element.get_or_add_pPr()
for old in list(title_ppr.findall(qn("w:pBdr"))):
    title_ppr.remove(old)
p_bdr = OxmlElement("w:pBdr")
for side in ("top", "left", "bottom", "right", "between", "bar"):
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "nil")
    p_bdr.append(edge)
title_ppr.append(p_bdr)

for qnum, qtype, stem, opts, answer in questions:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{qnum}. "); r.bold = True; r.font.size = Pt(11.5)
    add_marked(p, stem, bold=True, size=11)

    if qtype == "S":
        ip = doc.add_paragraph(); ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ip.paragraph_format.space_after = Pt(4); ip.paragraph_format.keep_with_next = True
        shape = ip.add_run().add_picture(str(qimg[qnum]), width=Inches(6.75 if opts is None else 6.55))
        shape._inline.docPr.set("descr", f"Structuurformule bij vraag {qnum}")
    if opts:
        for idx, (text, is_formula) in enumerate(opts):
            op = doc.add_paragraph()
            if idx < len(opts)-1: op.paragraph_format.keep_with_next = True
            add_option(op, "ABCD"[idx], text, is_formula)

# Document properties
doc.core_properties.title = "HAVO Scheikunde Namen en formules"
doc.core_properties.subject = "Meerkeuzequiz syllabus centraal examen 2027"
doc.core_properties.author = ""
doc.core_properties.keywords = "HAVO, scheikunde, molecuulformule, verhoudingsformule, structuurformule"

doc.save(OUT)
print(OUT)
